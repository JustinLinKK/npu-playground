"""Build the editable progress deck from the colocated, verified data snapshot.

Run: uv run --with python-pptx==1.0.2 python docs/presentations/npu-progress-2026-09-09/build_deck.py
"""

import json
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_MARK
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


HERE = Path(__file__).resolve().parent
DATA = json.loads((HERE / "evidence.json").read_text())
P = Presentation()
P.slide_width, P.slide_height = Inches(13.333333), Inches(7.5)
P.core_properties.title = "NPU translation: current progress and the hardware gap"
P.core_properties.subject = "Measured translation completion, offline correctness checks, and unmeasured edge-AI performance"
P.core_properties.author = "NPU Playground"
P.core_properties.keywords = "NPU, AMD, Intel, agentic translation, correctness, power, latency"
NAVY, INK, MUTED = "112337", "142B40", "546579"
TEAL, BLUE, AMBER, RED = "008778", "356BDC", "D68A19", "B34746"
WHITE, BG, LINE = "FFFFFF", "F5F7FA", "DFE6ED"
PALE_TEAL, PALE_BLUE, PALE_AMBER, PALE_RED = "E5F3F0", "EAF0FD", "FFF2DB", "FCECEC"
NOTES = []


def rgb(value):
    return RGBColor.from_string(value)


def box(slide, x, y, w, h, fill, border=None, rounded=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    for effect in shape._element.xpath("./p:style/a:effectRef"):
        effect.set("idx", "0")
    if border:
        shape.line.color.rgb = rgb(border)
        shape.line.width = Pt(0.8)
    else:
        shape.line.fill.background()
    if rounded:
        shape.adjustments[0] = 0.06
    return shape


def text(slide, x, y, w, h, value, size=20, color=INK, bold=False, align=None, link=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, line in enumerate(str(value).split("\n")):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.space_after = Pt(5)
        para.space_before = Pt(0)
        para.line_spacing = 1.08
        if align is not None:
            para.alignment = align
        run = para.add_run()
        run.text = line
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
        if link:
            run.hyperlink.address = link
    return shape


def slide(kicker, title, subtitle, source, dark=False):
    s = P.slides.add_slide(P.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(NAVY if dark else WHITE)
    text(s, 0.65, 0.30, 11.8, 0.23, kicker.upper(), 11, "5BD5C4" if dark else TEAL, True)
    text(s, 0.65, 0.77, 12, 0.68, title, 31, WHITE if dark else INK, True)
    if subtitle:
        text(s, 0.65, 1.47, 12, 0.52, subtitle, 17, "BCCCDD" if dark else MUTED)
    box(s, 0.65, 7.02, 12.02, 0.009, "3A4C61" if dark else LINE)
    text(s, 0.65, 7.13, 11.2, 0.18, source, 9, "BCCCDD" if dark else MUTED)
    text(s, 12.05, 7.11, 0.60, 0.22, f"{len(P.slides):02d}", 11, "BCCCDD" if dark else MUTED, align=PP_ALIGN.RIGHT)
    return s


def note(s, title, body, sources):
    content = body + "\n\nSources:\n" + "\n".join(sources)
    s.notes_slide.notes_text_frame.text = content
    NOTES.append(f"## Slide {len(P.slides)} — {title}\n\n{body}\n\nSources:\n\n" + "\n".join(f"- {v}" for v in sources))


def metric(s, x, y, w, value, label, detail, fill=BG, color=TEAL, h=1.42):
    box(s, x, y, w, h, fill, rounded=True)
    text(s, x + 0.20, y + 0.12, w - 0.40, 0.55, value, 32, color, True)
    text(s, x + 0.20, y + 0.73, w - 0.40, 0.32, label, 16, INK, True)
    text(s, x + 0.20, y + 1.10, w - 0.40, h - 1.10, detail, 12, MUTED)


E1 = "E1: runs/20260904T011639Z/{experiment.json, metrics.jsonl, summary.json}"
E2 = "E2: runs/container-upgrade-validation/{summary.json, corpus-final/report.json, fixtures/report.json}"
E3 = "E3: docs/validation-results.md and docs/container-simulation-vs-hardware-report.md"
PUBLIC = DATA["public_sources"]

# 1: Keep the two experiments and missing hardware evidence visibly distinct.
s = slide("NPU Playground / progress brief / 09 Sep 2026", "", "", "Sources: E1 historical experiment; E2 current evaluator validation. Full provenance in notes and evidence.json.", True)
text(s, 0.65, 0.98, 11.7, 1.65, "From LLM translation\nto validated NPU kernels", 40, WHITE, True)
text(s, 0.65, 2.65, 11.8, 0.47, "Agentic feedback improves completion; hardware testing is the next milestone.", 21, "BCCCDD")
cards = [
    (0.65, "01 / TRANSLATION", "70% → 95%", "Historical task completion", "One-shot → agentic\n14/20 → 19/20 cases", "7FA7FF"),
    (4.76, "02 / CORRECTNESS", "22/22", "Current offline checks pass", "Deterministic candidates\n264 host input cases", "5BD5C4"),
    (8.87, "03 / HARDWARE", "0 runs", "NPU execution blocked", "Latency and power\nare unmeasured", "F4BD69"),
]
for x, label, value, caption, detail, color in cards:
    box(s, x, 3.34, 3.81, 2.84, "1C334B", rounded=True)
    text(s, x + 0.23, 3.58, 3.35, 0.24, label, 11, color, True)
    text(s, x + 0.23, 4.05, 3.35, 0.67, value, 37, WHITE, True)
    text(s, x + 0.23, 4.91, 3.35, 0.50, caption, 17, WHITE, True)
    text(s, x + 0.23, 5.48, 3.35, 0.62, detail, 14, "BCCCDD")
text(s, 0.65, 6.43, 12, 0.35, "Different evidence sets: the 22-case evaluator suite is not a rerun of the historical LLM comparison.", 15, "BCCCDD")
note(s, "Progress at a glance", "The measured claim is that the evaluated one-shot baseline is unreliable, not that all current LLMs cannot translate NPU code. In the completed historical experiment, baseline completed 14 of 20 kernel-target tasks and agentic completed 19 of 20. AMD completion meant compilation only; Intel completion included historical CPU equivalence and NPU compilation. The newer 22-case suite verifies the upgraded evaluator using deterministic candidates and zero provider calls. It does not demonstrate a new 100% LLM translation rate. No exact NPU binary execution or NPU latency/power measurement is recorded.", [E1, E2, E3])

# 2: The actual baseline failures, without inventing model-wide conclusions.
s = slide("01 / Historical LLM experiment", "One-shot translation is not reliable enough", "10 kernels × 2 NPU targets; one model call and one compile attempt per task.", "Source: E1 • Completed 04 Sep 2026 UTC • Codex CLI 0.144.5 • Model setting recorded as ‘CLI default’.")
text(s, 0.65, 2.14, 5.1, 0.3, "KERNEL", 12, MUTED, True)
text(s, 5.70, 2.11, 1.15, 0.34, "AMD", 14, INK, True, PP_ALIGN.CENTER)
text(s, 7.00, 2.11, 1.15, 0.34, "Intel", 14, INK, True, PP_ALIGN.CENTER)
labels = {"cuda_reduce_sum": "CUDA · Reduce sum", "cuda_tiled_matmul": "CUDA · Tiled matrix multiply", "cuda_tiled_transpose": "CUDA · Tiled transpose", "cuda_vector_add": "CUDA · Vector add", "hip_gamma_correction": "HIP · Gamma correction", "hip_image_convolution": "HIP · Image convolution", "hip_moving_average": "HIP · Moving average", "triton_fused_attention": "Triton · Fused attention", "triton_fused_softmax": "Triton · Fused softmax", "triton_layer_norm": "Triton · Layer normalization"}
for i, (kernel, label) in enumerate(labels.items()):
    y = 2.56 + 0.346 * i
    box(s, 0.65, y - 0.035, 7.66, 0.336, BG if i % 2 == 0 else WHITE)
    text(s, 0.80, y, 4.65, 0.28, label, 14)
    for x, target in [(5.70, "amd_xdna2_npu2"), (7.00, "intel_npu_4000")]:
        passed = next(r["baseline_pass"] for r in DATA["historical"]["paired_cases"] if r["kernel"] == kernel and r["target_id"] == target)
        box(s, x, y - 0.005, 1.15, 0.28, PALE_TEAL if passed else PALE_RED, rounded=True)
        text(s, x, y + 0.012, 1.15, 0.23, "PASS" if passed else "FAIL", 11, TEAL if passed else RED, True, PP_ALIGN.CENTER)
metric(s, 8.80, 2.17, 3.87, "14/20", "tasks completed", "70% overall completion", PALE_BLUE, BLUE, 1.52)
metric(s, 8.80, 3.88, 3.87, "6 failures", "across both targets", "AMD: 4  |  Intel: 2", PALE_RED, RED, 1.52)
text(s, 8.80, 5.70, 3.87, 0.63, "This is one evaluated provider configuration, not a survey of all LLMs.", 15, MUTED)
text(s, 0.65, 6.44, 12, 0.35, "PASS criteria: AMD = compilation only; Intel = historical CPU equivalence + NPU compilation.", 14, MUTED)
note(s, "One-shot baseline", "The table is derived from metrics.jsonl and cross-checked against the completed experiment. AMD baseline passes: reduce sum, tiled matmul, tiled transpose, gamma correction, image convolution, fused softmax (6/10). Intel baseline fails image convolution and moving average (8/10 pass). PASS is historical task completion and must not be described as physical NPU correctness. The provider metadata names codex-cli 0.144.5 with model ‘CLI default’; a precise model version is not pinned in the experiment metadata. This limits broader claims about contemporary LLMs.", [E1])

# 3: A native PowerPoint chart, backed by an embedded editable workbook.
s = slide("02 / Historical LLM experiment", "Agentic feedback recovers 5 of 6 failures", "The framework analyzes, generates, compiles and repairs candidates using tool feedback.", "Source: E1 • Same 20 kernel-target tasks • Historical acceptance criteria • Search budgets differ.")
cd = CategoryChartData()
cd.categories = ["AMD", "Intel"]
groups = DATA["historical"]["groups"]
for method, label in [("baseline", "One-shot"), ("agentic", "Agentic")]:
    cd.add_series(label, [next(g["successes"] for g in groups if g["method"] == method and g["target_id"] == t) for t in ["amd_xdna2_npu2", "intel_npu_4000"]])
chart = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.65), Inches(2.06), Inches(7.55), Inches(3.28), cd).chart
chart.has_legend = True
chart.legend.position = XL_LEGEND_POSITION.BOTTOM
chart.legend.include_in_layout = False
chart.legend.font.name = "Arial"
chart.legend.font.size = Pt(13)
chart.chart_style = 10
chart.font.name = "Arial"
chart.font.size = Pt(13)
axis = chart.value_axis
axis.minimum_scale, axis.maximum_scale, axis.major_unit = 0, 10, 2
axis.tick_labels.font.size = Pt(11)
axis.tick_labels.font.color.rgb = rgb(MUTED)
axis.major_tick_mark = XL_TICK_MARK.NONE
axis.has_major_gridlines = True
axis.major_gridlines.format.line.color.rgb = rgb(LINE)
chart.category_axis.tick_labels.font.size = Pt(15)
chart.category_axis.major_tick_mark = XL_TICK_MARK.NONE
plot = chart.plots[0]
plot.gap_width = 95
plot.has_data_labels = True
plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
plot.data_labels.number_format = '0"/10"'
plot.data_labels.font.name = "Arial"
plot.data_labels.font.size = Pt(17)
plot.data_labels.font.bold = True
plot.data_labels.font.color.rgb = rgb(INK)
for series, color in zip(chart.series, [BLUE, TEAL]):
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = rgb(color)
    series.format.line.fill.background()
text(s, 0.90, 5.40, 7.05, 0.46, "AMD: compilation only   |   Intel: CPU equivalence + compilation", 12, MUTED, align=PP_ALIGN.CENTER)
metric(s, 8.80, 2.10, 3.87, "+25 pp", "completion improvement", "70% → 95%  |  14/20 → 19/20", PALE_TEAL, TEAL, 1.58)
text(s, 8.80, 4.04, 3.80, 0.35, "1 failure remains", 21, INK, True)
text(s, 8.80, 4.56, 3.80, 0.95, "AMD fused attention failed the intermediate-representation check.", 18, MUTED)
box(s, 0.65, 5.98, 12.02, 0.52, BG, rounded=True)
text(s, 0.83, 6.10, 11.66, 0.27, "SEARCH COST     Model calls: 20 → 127       Compile attempts: 20 → 216", 17, INK, True)
text(s, 0.65, 6.67, 12, 0.24, "This shows a benefit from the full agentic workflow; it is not an equal-compute comparison or a 95% device-correctness result.", 12, MUTED)
note(s, "Agentic results and cost", "The native chart uses final case completion, not the count of successful intermediate compile attempts. Agentic completes 9/10 AMD and 10/10 Intel tasks, compared with baseline 6/10 and 8/10. Five paired baseline failures become agentic successes; AMD fused attention remains unsuccessful with IR mismatch {'output': 0.5390625} after three analysis attempts. Agentic uses 127 provider calls and 216 compile attempts versus 20 and 20 for baseline. Search configuration: three rounds, branching factor three, up to two repair retries per candidate. Therefore the observed +25 percentage points is a workflow result at different budgets, not an isolated causal estimate of agent architecture. Intel host equivalence uses the historical evaluator, not the upgraded independent 12-input suite. AMD historical results contain no source or device correctness proof.", [E1])

# 4: Evidence that the upgraded evaluator enforces scoped correctness.
s = slide("03 / Current evaluator validation", "The framework now checks correctness explicitly", "Independent references, host execution and AMD dataflow checks form the offline acceptance gate.", "Sources: E2, E3 • 11 deterministic candidates per target • 12 inputs per candidate • No LLM calls in this suite.")
for i, (label, detail) in enumerate([("01  REFERENCE", "Independent expected outputs"), ("02  EXECUTE", "Candidate source / CPU graph"), ("03  CHECK", "Numerics + AMD dataflow"), ("04  COMPILE", "Pinned NPU toolchain")]):
    x = 0.65 + 3.08 * i
    box(s, x, 2.02, 2.78, 0.79, BG, rounded=True)
    text(s, x + 0.13, 2.15, 2.52, 0.22, label, 12, TEAL, True)
    text(s, x + 0.13, 2.51, 2.52, 0.24, detail, 11, MUTED)
    if i < 3:
        text(s, x + 2.80, 2.25, 0.27, 0.33, "→", 19, MUTED, align=PP_ALIGN.CENTER)
for x, value, label, detail in [(0.65, "22/22", "offline cases pass", "11 AMD + 11 Intel candidates"), (4.76, "264", "host input cases checked", "12 inputs × 22 candidates"), (8.87, "396", "AMD case / schedule checks", "132 inputs × 3 schedules")]:
    metric(s, x, 3.08, 3.81, value, label, detail, PALE_TEAL, TEAL, 1.43)
rows = [("Current validation gate", "AMD", "Intel"), ("Host output + target compilation", "11/11", "11/11"), ("FIFO / DMA dataflow simulation", "11/11", "Not implemented")]
for i, row in enumerate(rows):
    y = 4.76 + i * 0.36
    box(s, 0.65, y, 12.02, 0.35, BG if i % 2 == 0 else WHITE)
    for x, w, label in zip([0.82, 7.65, 10.1], [6.55, 1.75, 2.38], row):
        text(s, x, y + 0.065, w, 0.27, label, 13, MUTED if i == 0 else INK, i == 0)
box(s, 0.65, 6.00, 12.02, 0.42, PALE_RED, rounded=True)
text(s, 0.82, 6.09, 11.68, 0.27, "Faults caught: wrong arithmetic → mismatch   |   bad DMA stride → bounds error   |   missing FIFO release → deadlock", 12, RED, True)
text(s, 0.65, 6.63, 12, 0.28, "These deterministic cases validate the evaluator; the historical LLM experiment has not been rerun under this stricter policy.", 13, MUTED)
note(s, "Correctness acceptance gates", "The upgraded suite contains 11 supported static contracts per target: the ten classic manifests plus sigmoid. Each candidate runs on five normal seeds plus seven boundary input families. Across the targets there are 264 host input cases. AMD executes 132 unique dataflow input cases under three schedules, yielding 396 case/schedule comparisons; those are not 396 independent inputs. All 22 deterministic candidate-target cases pass the offline policy; eight additional positive fixtures also pass. Negative checks exercise arithmetic mutation, DMA bounds, missing FIFO release/deadlock and other boundaries. This supports scoped source/graph and modeled-dataflow correctness, not a proof for arbitrary kernels, original CUDA/HIP/Triton execution, or final device binaries. AMD host and dataflow checks share a topology model. The current suite has zero provider calls and is separate from the historical LLM experiment. Current exact-binary target execution remains blocked on both vendors.", [E2, E3])

# 5: The vendor distinction is explicit rather than hidden in a footnote.
s = slide("04 / AMD and Intel scenarios", "AMD and Intel validate different representations", "Both paths can check supported computations and build target artifacts inside containers.", "Sources: E2, E3; P1 AMD dataflow pre-simulation; P2 Intel OpenVINO NPU documentation. Links in notes.")
for x, vendor, device, color, fill, blocks in [
    (0.65, "AMD", "XDNA2 / npu2 / AIE2P", TEAL, PALE_TEAL, [("EXECUTES ON HOST", "Candidate C++ through audited host support; outputs checked against a reference."), ("MODELS DATA MOVEMENT", "Bounded MLIR FIFO / DMA model checks transfers, bounds and deadlocks."), ("COMPILES FOR NPU", "MLIR-AIE / Peano produce target artifacts, including xclbin and instructions.")]),
    (6.81, "Intel", "OpenVINO / NPU 4000", BLUE, PALE_BLUE, [("EXECUTES ON CPU", "Candidate OpenVINO graph runs on the CPU backend; outputs checked against a reference."), ("CHECKS GRAPH SUPPORT", "Offline NPU compilation tests whether the pinned compiler accepts the graph."), ("COMPILES FOR NPU", "The compiler exports an NPU blob. CPU execution does not execute that blob.")]),
]:
    box(s, x, 2.08, 5.86, 4.24, BG, rounded=True)
    box(s, x, 2.08, 5.86, 0.83, fill, rounded=True)
    text(s, x + 0.23, 2.24, 1.5, 0.4, vendor, 25, color, True)
    text(s, x + 1.9, 2.36, 3.69, 0.3, device, 14, MUTED, align=PP_ALIGN.RIGHT)
    for j, (label, detail) in enumerate(blocks):
        y = 3.16 + 1.02 * j
        text(s, x + 0.23, y, 5.4, 0.24, label, 11, color, True)
        text(s, x + 0.23, y + 0.33, 5.32, 0.57, detail, 16, INK)
box(s, 0.65, 6.53, 12.02, 0.35, PALE_AMBER, rounded=True)
text(s, 0.82, 6.59, 11.7, 0.25, "Shared gap: 0 exact target-binary executions; all 22 corpus cases report no validated target executor.", 13, "82510A", True)
note(s, "Vendor-specific evidence", "AMD: the upgraded path executes supported candidate C++ through host compatibility support, models emitted MLIR topology and bounded FIFO/DMA behavior, and compiles with the pinned MLIR-AIE/Peano toolchain. Unsupported intrinsics or topology are rejected or marked unsupported rather than treated as verified. This is source/dataflow validation, not an AIE2P instruction-set or cycle-accurate binary simulator. Intel: a candidate graph executes through OpenVINO CPU and also undergoes offline NPU compilation. CPU graph correctness is separate from correctness of the compiled NPU blob. Both profiles have no hardware runner, and every target_execution stage is blocked with NO_VALIDATED_TARGET_EXECUTOR; target-executed validation returns exit code 3. A container can host a real device runner when compatible hardware and drivers are passed through, but these current offline paths do not provide that execution. Vendor documentation is supporting context; the counts and executor status come from local reports.", [E2, E3, f"P1: {PUBLIC['P1']['url']}", f"P2: {PUBLIC['P2']['url']}"])

# 6: Unknown performance is shown as unknown, never a zero value.
s = slide("05 / Why hardware matters for edge AI", "A correct kernel still needs a hardware benchmark", "Edge workloads need acceptable latency and energy use as well as correct outputs.", "Sources: E2, E3; P3 MLPerf Tiny (quality, latency, optional energy); P4 MLPerf Edge power methodology.", True)
items = [
    (0.65, "CORRECTNESS", "Offline checked", "Does it compute the right result?", "Source / graph evidence exists.\nExact binary still needs NPU testing.", "5BD5C4"),
    (4.76, "PERFORMANCE", "Unmeasured", "Does it meet the response budget?", "Measure NPU p50 / p95 latency\nand completed work per second.", "F4BD69"),
    (8.87, "POWER + ENERGY", "Unmeasured", "Does it fit the device’s energy budget?", "Measure average watts and\nmillijoules per completed inference.", "F4BD69"),
]
for x, label, state, question, detail, color in items:
    box(s, x, 2.20, 3.81, 3.30, "1C334B", rounded=True)
    text(s, x + 0.22, 2.47, 3.37, 0.3, label, 12, color, True)
    text(s, x + 0.22, 3.04, 3.37, 0.52, state, 28, WHITE, True)
    text(s, x + 0.22, 3.84, 3.33, 0.75, question, 20, WHITE)
    text(s, x + 0.22, 4.84, 3.37, 0.59, detail, 14, "BCCCDD")
text(s, 0.65, 5.89, 12, 0.44, "Container checks cannot establish NPU speed, power or energy per inference.", 20, WHITE, True)
text(s, 0.65, 6.49, 12, 0.34, "Measure on AMD and Intel hardware; state whether power covers the NPU, the board or the whole system.", 15, "BCCCDD")
note(s, "Performance and power remain open", "No NPU latency, throughput, power or energy result is available from the current validation. Missing latency is null, not zero; power and energy were not collected. Functional simulations do not model the device’s actual clock behavior, runtime overhead, memory traffic timing, thermal state or physical power draw with validated accuracy. Hardware measurements should report warmup policy, repeated timings (including p50 and p95), throughput, average power and energy per completed inference. Energy per inference equals measured energy over the measurement window divided by completed inferences. Keep the workload, precision, shape, batch, transfer boundary and measurement boundary explicit. A power rating or CPU runtime is not a substitute for a measured device workload. MLPerf Tiny includes quality, latency and optional energy; MLPerf Edge explains whole-system power measurement. These sources motivate the proposed measurements; our project has no MLPerf submission or performance result.", [E2, E3, f"P3: {PUBLIC['P3']['url']}", f"P4: {PUBLIC['P4']['url']}"])

# 7: Concrete next work, clearly labelled as planned.
s = slide("06 / Next milestones — planned", "Close the loop on real AMD and Intel hardware", "The next result should connect translation success, device correctness, latency and energy.", "Proposed work based on the gaps in E1–E3. No new hardware or LLM experiments were run to prepare this deck.")
steps = [
    ("01", "Rerun the translation comparison", "Use the upgraded correctness policy for both methods.", "Pin the model and compare matched tasks and budgets."),
    ("02", "Execute the exact target artifacts", "Connect each NPU, its drivers and a validated device runner.", "Compare device outputs with the same independent references."),
    ("03", "Measure and optimize for the edge", "Collect latency, throughput, watts and energy per inference.", "Compare only candidates that pass device correctness checks."),
]
for i, (num, title, body, detail) in enumerate(steps):
    y = 2.13 + i * 1.24
    box(s, 0.65, y, 0.66, 0.66, PALE_TEAL, rounded=True)
    text(s, 0.65, y + 0.13, 0.66, 0.34, num, 20, TEAL, True, PP_ALIGN.CENTER)
    text(s, 1.61, y - 0.015, 10.7, 0.42, title, 22, INK, True)
    text(s, 1.61, y + 0.47, 10.7, 0.30, body, 17, INK)
    text(s, 1.61, y + 0.83, 10.7, 0.30, detail, 15, MUTED)
box(s, 0.65, 6.21, 12.02, 0.58, PALE_BLUE, rounded=True)
text(s, 0.87, 6.36, 11.58, 0.32, "Deliverable: a correctness-filtered latency and energy comparison for both NPU targets.", 18, BLUE, True)
note(s, "Next milestones", "These are proposed next steps, not completed work. First rerun the baseline and agentic translation experiment under the new offline acceptance policy, with an explicitly pinned model and a matched-budget comparison in addition to the existing workflow comparison. Second validate a device runner on each target and execute the exact compiled artifacts, verifying outputs before collecting performance. Third benchmark fixed workloads and precision under documented warmup and steady-state conditions, measuring latency, throughput, power and energy. State whether host transfers, loading and compilation are included. A useful result is a set of candidates that preserve device correctness while improving measured latency or energy, with the measurement boundaries disclosed.", [E1, E2, E3])

# 8: Shareable provenance and limitations, with the full data in the sidecar.
s = slide("Appendix / Evidence and interpretation", "What the numbers prove — and what remains open", "All charts and counts are backed by the saved project reports; source hashes are in evidence.json.", "Project: npu-playground • Snapshot prepared 09 Sep 2026 • Speaker notes contain source paths, links and metric definitions.")
box(s, 0.65, 2.13, 5.86, 3.46, BG, rounded=True)
text(s, 0.88, 2.36, 5.4, 0.35, "E1 / Historical translation experiment", 19, BLUE, True)
text(s, 0.88, 2.95, 5.27, 2.10, "40 recorded cases: 20 per method\n10 kernels × AMD and Intel\nRevision: 74f8ae752b23\nCodex CLI 0.144.5; model: CLI default\nAMD compile-only; Intel CPU + compile\nUnequal search budgets", 16, INK)
text(s, 0.88, 5.20, 5.24, 0.30, "Supports workflow completion, not device correctness.", 13, MUTED)
box(s, 6.81, 2.13, 5.86, 3.46, BG, rounded=True)
text(s, 7.04, 2.36, 5.4, 0.35, "E2 / Current evaluator validation", 19, TEAL, True)
text(s, 7.04, 2.95, 5.27, 2.10, "22 deterministic candidate-target cases\n12 inputs per candidate; 3 AMD schedules\nBase: d3efa60a08e5 + working-tree upgrade\n0 provider calls; 0 target binary executions\nIndependent references and negative checks\nNot a rerun of the historical LLM experiment", 16, INK)
text(s, 7.04, 5.20, 5.24, 0.30, "Supports scoped offline checks, not exhaustive proof.", 13, MUTED)
text(s, 0.65, 5.91, 12, 0.3, "PUBLIC CONTEXT — CLICKABLE SOURCES", 11, MUTED, True)
for x, y, key, label in [(0.65, 6.29, "P1", "P1  AMD dataflow pre-simulation"), (6.81, 6.29, "P2", "P2  Intel OpenVINO NPU device"), (0.65, 6.65, "P3", "P3  MLPerf Tiny: quality, latency, energy"), (6.81, 6.65, "P4", "P4  MLPerf Edge: power measurement")]:
    text(s, x, y, 5.8, 0.27, label, 13, BLUE, link=PUBLIC[key]["url"])
note(s, "Evidence and interpretation", "Evidence snapshot: 2026-09-09. Historical experiment ID 20260904T011639Z completed at 2026-09-04T06:42:44 UTC. The historical revision is 74f8ae752b2396cfb3f5e44226df93f2d8eb277d. The upgraded evaluator was validated on base d3efa60a08e5807c54b40832f3dcb2549f44ea00 plus uncommitted implementation changes. Image identities and compiler fingerprints, original per-case metrics, aggregate counts and SHA-256 hashes of source reports are preserved in the colocated evidence.json. Do not compare 19/20 and 22/22 as a before/after translation accuracy experiment: the candidates, evaluation criteria and use of LLMs differ. The presentation wording intentionally limits the first claim to the measured baseline and limits correctness to the representations actually checked. Full test run before the final alias guard: 98 passed including 23 container tests. Final CPU run: 76 passed, 23 deselected; final boundary checks: 47 passed. The corpus and fixtures were rerun after the guard. Those software checks do not provide hardware performance evidence.", [E1, E2, E3] + [f"{key}: {item['title']} — {item['url']}" for key, item in PUBLIC.items()])

P.save(HERE / "npu-progress.pptx")
(HERE / "speaker-notes.md").write_text("# NPU translation progress — speaker notes\n\n" + "\n\n".join(NOTES) + "\n")
assert len(P.slides) == 8
for i, s in enumerate(P.slides, 1):
    for shape in s.shapes:
        assert shape.left >= 0 and shape.top >= 0, (i, shape.name)
        assert shape.left + shape.width <= P.slide_width, (i, shape.name)
        assert shape.top + shape.height <= P.slide_height, (i, shape.name)
print(f"Built {len(P.slides)} editable slides: {HERE / 'npu-progress.pptx'}")
